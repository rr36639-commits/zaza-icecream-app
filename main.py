import os
import io
import json
import csv
import secrets
import hashlib
import hmac
import base64
import logging
from datetime import datetime, date, timedelta, timezone
from decimal import Decimal, InvalidOperation
from typing import Optional, Any

import requests
from fastapi import FastAPI, Request, UploadFile, File, Form, HTTPException, Depends
from fastapi.responses import HTMLResponse, JSONResponse, StreamingResponse, RedirectResponse
from fastapi.middleware.cors import CORSMiddleware
from starlette.middleware.sessions import SessionMiddleware

from sqlalchemy import (
    create_engine,
    Column,
    Integer,
    String,
    Text,
    Float,
    DateTime,
    Boolean,
    ForeignKey,
    func,
    text,
)
from sqlalchemy.orm import declarative_base, sessionmaker, Session, relationship

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

try:
    from google import genai
    from google.genai import types
except Exception:
    genai = None
    types = None


# -----------------------------------------------------------------------------
# Configuration
# -----------------------------------------------------------------------------

APP_NAME = "رمضان جاد / ظاظا - موزع الآيس كريم"
DATABASE_URL = os.getenv("DATABASE_URL", "sqlite:///./ramadan_icecream.db")
SECRET_KEY = os.getenv("SECRET_KEY", "CHANGE_THIS_SECRET_KEY_IN_PRODUCTION")
ADMIN_USERNAME = os.getenv("ADMIN_USERNAME", "admin")
ADMIN_PASSWORD = os.getenv("ADMIN_PASSWORD", "admin123")
GEMINI_API_KEY = os.getenv("GEMINI_API_KEY", "")
GEMINI_MODEL = os.getenv("GEMINI_MODEL", "gemini-1.5-flash")
TELEGRAM_BOT_TOKEN = os.getenv("TELEGRAM_BOT_TOKEN", "")
TELEGRAM_CHAT_ID = os.getenv("TELEGRAM_CHAT_ID", "")
FACEBOOK_VERIFY_TOKEN = os.getenv("FACEBOOK_VERIFY_TOKEN", "")
WHATSAPP_VERIFY_TOKEN = os.getenv("WHATSAPP_VERIFY_TOKEN", "")
MAX_UPLOAD_MB = int(os.getenv("MAX_UPLOAD_MB", "10"))

logging.basicConfig(level=logging.INFO)
logger = logging.getLogger("ramadan_icecream")

if DATABASE_URL.startswith("postgres://"):
    DATABASE_URL = DATABASE_URL.replace("postgres://", "postgresql://", 1)

connect_args = {"check_same_thread": False} if DATABASE_URL.startswith("sqlite") else {}
engine = create_engine(
    DATABASE_URL,
    connect_args=connect_args,
    pool_pre_ping=True,
)
SessionLocal = sessionmaker(bind=engine, autoflush=False, autocommit=False)
Base = declarative_base()


# -----------------------------------------------------------------------------
# Database models
# -----------------------------------------------------------------------------

class User(Base):
    __tablename__ = "users"

    id = Column(Integer, primary_key=True)
    username = Column(String(120), unique=True, nullable=False, index=True)
    password_hash = Column(String(255), nullable=False)
    role = Column(String(50), nullable=False, default="admin")
    is_active = Column(Boolean, nullable=False, default=True)
    created_at = Column(DateTime, nullable=False, default=datetime.utcnow)


class Invoice(Base):
    __tablename__ = "invoices"

    id = Column(Integer, primary_key=True)
    customer_name = Column(String(255), nullable=False, default="عميل نقدي")
    items_json = Column(Text, nullable=False, default="[]")
    total_amount = Column(Float, nullable=False, default=0)
    created_at = Column(DateTime, nullable=False, default=datetime.utcnow, index=True)


class IntegrationConfig(Base):
    __tablename__ = "integration_configs"

    id = Column(Integer, primary_key=True)
    platform = Column(String(80), unique=True, nullable=False, index=True)
    config_json = Column(Text, nullable=False, default="{}")


class Product(Base):
    __tablename__ = "products"

    id = Column(Integer, primary_key=True)
    name = Column(String(255), nullable=False, index=True)
    code = Column(String(100), unique=True, nullable=True, index=True)
    category = Column(String(120), nullable=True)
    unit = Column(String(50), nullable=False, default="قطعة")
    purchase_price = Column(Float, nullable=False, default=0)
    sale_price = Column(Float, nullable=False, default=0)
    quantity = Column(Float, nullable=False, default=0)
    min_quantity = Column(Float, nullable=False, default=0)
    image_url = Column(String(500), nullable=True)
    expiry_date = Column(String(30), nullable=True)
    created_at = Column(DateTime, nullable=False, default=datetime.utcnow)


class Vehicle(Base):
    __tablename__ = "vehicles"

    id = Column(Integer, primary_key=True)
    name = Column(String(120), unique=True, nullable=False)
    color = Column(String(50), nullable=True)
    driver_name = Column(String(150), nullable=True)
    is_active = Column(Boolean, nullable=False, default=True)
    created_at = Column(DateTime, nullable=False, default=datetime.utcnow)


class VehicleInventory(Base):
    __tablename__ = "vehicle_inventory"

    id = Column(Integer, primary_key=True)
    vehicle_id = Column(Integer, ForeignKey("vehicles.id"), nullable=False, index=True)
    product_id = Column(Integer, ForeignKey("products.id"), nullable=False, index=True)
    quantity = Column(Float, nullable=False, default=0)

    vehicle = relationship("Vehicle")
    product = relationship("Product")


class StockMovement(Base):
    __tablename__ = "stock_movements"

    id = Column(Integer, primary_key=True)
    product_id = Column(Integer, ForeignKey("products.id"), nullable=False, index=True)
    vehicle_id = Column(Integer, ForeignKey("vehicles.id"), nullable=True, index=True)
    movement_type = Column(String(60), nullable=False)
    quantity = Column(Float, nullable=False)
    reference_id = Column(String(100), nullable=True)
    note = Column(Text, nullable=True)
    created_by = Column(String(120), nullable=True)
    created_at = Column(DateTime, nullable=False, default=datetime.utcnow, index=True)


class Sale(Base):
    __tablename__ = "sales"

    id = Column(Integer, primary_key=True)
    vehicle_id = Column(Integer, ForeignKey("vehicles.id"), nullable=True, index=True)
    customer_id = Column(Integer, nullable=True, index=True)
    invoice_id = Column(Integer, ForeignKey("invoices.id"), nullable=True, index=True)
    payment_method = Column(String(50), nullable=False, default="نقدي")
    subtotal = Column(Float, nullable=False, default=0)
    discount = Column(Float, nullable=False, default=0)
    total = Column(Float, nullable=False, default=0)
    created_by = Column(String(120), nullable=True)
    created_at = Column(DateTime, nullable=False, default=datetime.utcnow, index=True)


class SaleItem(Base):
    __tablename__ = "sale_items"

    id = Column(Integer, primary_key=True)
    sale_id = Column(Integer, ForeignKey("sales.id"), nullable=False, index=True)
    product_id = Column(Integer, ForeignKey("products.id"), nullable=False, index=True)
    quantity = Column(Float, nullable=False)
    unit_price = Column(Float, nullable=False)
    purchase_price = Column(Float, nullable=False, default=0)
    total = Column(Float, nullable=False, default=0)


class Customer(Base):
    __tablename__ = "customers"

    id = Column(Integer, primary_key=True)
    name = Column(String(255), nullable=False, index=True)
    phone = Column(String(80), nullable=True)
    area = Column(String(150), nullable=True)
    customer_type = Column(String(80), nullable=True)
    balance = Column(Float, nullable=False, default=0)
    total_purchases = Column(Float, nullable=False, default=0)
    created_at = Column(DateTime, nullable=False, default=datetime.utcnow)


class CustomerTransaction(Base):
    __tablename__ = "customer_transactions"

    id = Column(Integer, primary_key=True)
    customer_id = Column(Integer, ForeignKey("customers.id"), nullable=False, index=True)
    transaction_type = Column(String(50), nullable=False)
    amount = Column(Float, nullable=False)
    reference_id = Column(String(100), nullable=True)
    note = Column(Text, nullable=True)
    created_at = Column(DateTime, nullable=False, default=datetime.utcnow)


class ReturnRecord(Base):
    __tablename__ = "returns"

    id = Column(Integer, primary_key=True)
    vehicle_id = Column(Integer, ForeignKey("vehicles.id"), nullable=True)
    product_id = Column(Integer, ForeignKey("products.id"), nullable=False)
    quantity = Column(Float, nullable=False)
    reason = Column(String(255), nullable=True)
    created_by = Column(String(120), nullable=True)
    created_at = Column(DateTime, nullable=False, default=datetime.utcnow)


class Damage(Base):
    __tablename__ = "damages"

    id = Column(Integer, primary_key=True)
    vehicle_id = Column(Integer, ForeignKey("vehicles.id"), nullable=True)
    product_id = Column(Integer, ForeignKey("products.id"), nullable=False)
    quantity = Column(Float, nullable=False)
    reason = Column(String(255), nullable=True)
    created_by = Column(String(120), nullable=True)
    created_at = Column(DateTime, nullable=False, default=datetime.utcnow)


class Expense(Base):
    __tablename__ = "expenses"

    id = Column(Integer, primary_key=True)
    vehicle_id = Column(Integer, ForeignKey("vehicles.id"), nullable=True)
    amount = Column(Float, nullable=False)
    reason = Column(String(255), nullable=False)
    created_by = Column(String(120), nullable=True)
    created_at = Column(DateTime, nullable=False, default=datetime.utcnow, index=True)


class DailySettlement(Base):
    __tablename__ = "daily_settlements"

    id = Column(Integer, primary_key=True)
    vehicle_id = Column(Integer, ForeignKey("vehicles.id"), nullable=False)
    settlement_date = Column(String(20), nullable=False, index=True)
    received_value = Column(Float, nullable=False, default=0)
    sales_value = Column(Float, nullable=False, default=0)
    returns_value = Column(Float, nullable=False, default=0)
    damage_value = Column(Float, nullable=False, default=0)
    expenses_value = Column(Float, nullable=False, default=0)
    expected_cash = Column(Float, nullable=False, default=0)
    actual_cash = Column(Float, nullable=False, default=0)
    difference = Column(Float, nullable=False, default=0)
    closed_by = Column(String(120), nullable=True)
    created_at = Column(DateTime, nullable=False, default=datetime.utcnow)


class Payment(Base):
    __tablename__ = "payments"

    id = Column(Integer, primary_key=True)
    customer_id = Column(Integer, ForeignKey("customers.id"), nullable=True)
    vehicle_id = Column(Integer, ForeignKey("vehicles.id"), nullable=True)
    amount = Column(Float, nullable=False)
    payment_method = Column(String(50), nullable=False, default="نقدي")
    note = Column(Text, nullable=True)
    created_by = Column(String(120), nullable=True)
    created_at = Column(DateTime, nullable=False, default=datetime.utcnow)


class Notification(Base):
    __tablename__ = "notifications"

    id = Column(Integer, primary_key=True)
    title = Column(String(255), nullable=False)
    body = Column(Text, nullable=False)
    level = Column(String(30), nullable=False, default="info")
    is_read = Column(Boolean, nullable=False, default=False)
    created_at = Column(DateTime, nullable=False, default=datetime.utcnow)


class AuditLog(Base):
    __tablename__ = "audit_logs"

    id = Column(Integer, primary_key=True)
    username = Column(String(120), nullable=True)
    action = Column(String(120), nullable=False)
    entity = Column(String(120), nullable=True)
    entity_id = Column(String(100), nullable=True)
    old_data = Column(Text, nullable=True)
    new_data = Column(Text, nullable=True)
    created_at = Column(DateTime, nullable=False, default=datetime.utcnow, index=True)


# -----------------------------------------------------------------------------
# Database bootstrap / self-healing
# -----------------------------------------------------------------------------

def hash_password(password: str) -> str:
    salt = secrets.token_bytes(16)
    digest = hashlib.pbkdf2_hmac("sha256", password.encode("utf-8"), salt, 210000)
    return base64.b64encode(salt + digest).decode("ascii")


def verify_password(password: str, encoded: str) -> bool:
    try:
        raw = base64.b64decode(encoded.encode("ascii"))
        salt = raw[:16]
        expected = raw[16:]
        actual = hashlib.pbkdf2_hmac("sha256", password.encode("utf-8"), salt, 210000)
        return hmac.compare_digest(actual, expected)
    except Exception:
        return False


def initialize_database() -> None:
    try:
        Base.metadata.create_all(bind=engine)
        db = SessionLocal()

        admin = db.query(User).filter(User.username == ADMIN_USERNAME).first()
        if not admin:
            admin = User(
                username=ADMIN_USERNAME,
                password_hash=hash_password(ADMIN_PASSWORD),
                role="admin",
                is_active=True,
            )
            db.add(admin)

        existing_vehicles = {v.name for v in db.query(Vehicle).all()}
        if "العربية الحمراء" not in existing_vehicles:
            db.add(Vehicle(name="العربية الحمراء", color="red", driver_name=""))
        if "العربية البيضاء" not in existing_vehicles:
            db.add(Vehicle(name="العربية البيضاء", color="white", driver_name=""))

        db.commit()
        db.close()
        logger.info("Database initialized successfully.")
    except Exception:
        logger.exception("Database initialization failed. Application will continue.")


initialize_database()


# -----------------------------------------------------------------------------
# FastAPI application
# -----------------------------------------------------------------------------

app = FastAPI(
    title=APP_NAME,
    version="1.0.0",
    docs_url="/docs",
    redoc_url="/redoc",
)

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

app.add_middleware(
    SessionMiddleware,
    secret_key=SECRET_KEY,
    session_cookie="ramadan_admin_session",
    max_age=60 * 60 * 24 * 7,
    same_site="lax",
    https_only=False,
)


def get_db():
    db = SessionLocal()
    try:
        yield db
    finally:
        db.close()


def current_user(request: Request, db: Session) -> Optional[User]:
    username = request.session.get("username")
    if not username:
        return None
    user = db.query(User).filter(User.username == username, User.is_active.is_(True)).first()
    return user


def require_user(request: Request, db: Session = Depends(get_db)) -> User:
    user = current_user(request, db)
    if not user:
        raise HTTPException(status_code=401, detail="يجب تسجيل الدخول")
    return user


def require_admin(request: Request, db: Session = Depends(get_db)) -> User:
    user = require_user(request, db)
    if user.role != "admin":
        raise HTTPException(status_code=403, detail="هذه العملية للمدير فقط")
    return user


def audit(
    db: Session,
    username: str,
    action: str,
    entity: str = "",
    entity_id: str = "",
    old_data: Any = None,
    new_data: Any = None,
):
    db.add(
        AuditLog(
            username=username,
            action=action,
            entity=entity,
            entity_id=str(entity_id or ""),
            old_data=json.dumps(old_data, ensure_ascii=False, default=str) if old_data is not None else None,
            new_data=json.dumps(new_data, ensure_ascii=False, default=str) if new_data is not None else None,
        )
    )


def safe_float(value: Any, default: float = 0.0) -> float:
    try:
        if value is None or value == "":
            return default
        return float(value)
    except (ValueError, TypeError):
        return default


def json_loads_safe(value: str, default: Any):
    try:
        return json.loads(value)
    except Exception:
        return default


def serialize_invoice(invoice: Invoice) -> dict:
    return {
        "id": invoice.id,
        "invoice_number": invoice.id,
        "customer_name": invoice.customer_name,
        "items": json_loads_safe(invoice.items_json, []),
        "total_amount": invoice.total_amount,
        "created_at": invoice.created_at.isoformat() if invoice.created_at else None,
    }


def telegram_config(db: Session) -> dict:
    config = db.query(IntegrationConfig).filter(IntegrationConfig.platform == "telegram").first()
    if not config:
        return {}
    return json_loads_safe(config.config_json, {})


def send_telegram_message(text_message: str, db: Session) -> bool:
    config = telegram_config(db)
    token = config.get("bot_token") or TELEGRAM_BOT_TOKEN
    chat_id = config.get("chat_id") or TELEGRAM_CHAT_ID

    if not token or not chat_id:
        return False

    try:
        response = requests.post(
            f"https://api.telegram.org/bot{token}/sendMessage",
            json={"chat_id": chat_id, "text": text_message},
            timeout=15,
        )
        response.raise_for_status()
        return True
    except Exception:
        logger.exception("Telegram notification failed.")
        return False


# -----------------------------------------------------------------------------
# Authentication
# -----------------------------------------------------------------------------

@app.get("/login", response_class=HTMLResponse)
def login_page(request: Request):
    if request.session.get("username"):
        return RedirectResponse("/", status_code=303)

    return HTMLResponse(
        """
        <!doctype html>
        <html lang="ar" dir="rtl">
        <head>
            <meta charset="utf-8">
            <meta name="viewport" content="width=device-width,initial-scale=1">
            <title>تسجيل الدخول</title>
            <style>
                *{box-sizing:border-box}
                body{margin:0;font-family:Arial,sans-serif;background:#f5f7fb;display:flex;min-height:100vh;align-items:center;justify-content:center}
                .card{width:min(420px,92vw);background:#fff;border-radius:24px;padding:28px;box-shadow:0 15px 50px rgba(0,0,0,.1)}
                h1{margin:0 0 8px}
                p{color:#667085}
                input{width:100%;padding:14px;margin:8px 0;border:1px solid #ddd;border-radius:12px;font-size:16px}
                button{width:100%;padding:14px;border:0;border-radius:12px;background:#111827;color:white;font-size:17px;font-weight:bold;margin-top:10px}
                .error{color:#b42318;background:#fef3f2;padding:10px;border-radius:10px;margin-bottom:10px}
            </style>
        </head>
        <body>
            <div class="card">
                <h1>🍦 رمضان جاد / ظاظا</h1>
                <p>إدارة توزيع وبيع الآيس كريم</p>
                <div id="error" class="error" style="display:none"></div>
                <form method="post" action="/auth/login">
                    <input name="username" placeholder="اسم المستخدم" required autocomplete="username">
                    <input name="password" type="password" placeholder="كلمة المرور" required autocomplete="current-password">
                    <button type="submit">دخول المدير</button>
                </form>
            </div>
        </body>
        </html>
        """
    )


@app.post("/auth/login")
def auth_login(
    request: Request,
    username: str = Form(...),
    password: str = Form(...),
    db: Session = Depends(get_db),
):
    user = db.query(User).filter(User.username == username, User.is_active.is_(True)).first()
    if not user or not verify_password(password, user.password_hash):
        return HTMLResponse(
            "<h3 style='font-family:Arial;text-align:center;margin-top:50px'>بيانات الدخول غير صحيحة.<br><a href='/login'>العودة</a></h3>",
            status_code=401,
        )

    request.session.clear()
    request.session["username"] = user.username
    request.session["role"] = user.role
    return RedirectResponse("/", status_code=303)


@app.post("/logout")
def logout(request: Request):
    request.session.clear()
    return RedirectResponse("/login", status_code=303)


# -----------------------------------------------------------------------------
# Dashboard / SPA
# -----------------------------------------------------------------------------

@app.get("/", response_class=HTMLResponse)
def home(request: Request, db: Session = Depends(get_db)):
    user = current_user(request, db)
    if not user:
        return RedirectResponse("/login", status_code=303)

    html = r"""
<!doctype html>
<html lang="ar" dir="rtl">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width,initial-scale=1,viewport-fit=cover">
<meta name="theme-color" content="#111827">
<title>رمضان جاد / ظاظا</title>
<style>
:root{--bg:#f5f7fb;--card:#fff;--text:#111827;--muted:#667085;--primary:#111827;--border:#e5e7eb;--green:#067647;--red:#b42318;--blue:#175cd3;--shadow:0 8px 30px rgba(16,24,40,.07)}
*{box-sizing:border-box}
body{margin:0;background:var(--bg);color:var(--text);font-family:Arial,Tahoma,sans-serif;padding-bottom:82px}
button,input,select,textarea{font:inherit}
button{cursor:pointer}
.top{position:sticky;top:0;z-index:10;background:#fff;border-bottom:1px solid var(--border);padding:13px 16px;display:flex;align-items:center;justify-content:space-between}
.brand{font-weight:800}.user{font-size:12px;color:var(--muted)}
main{max-width:1100px;margin:auto;padding:16px}
.page{display:none}.page.active{display:block}
.grid{display:grid;grid-template-columns:repeat(2,minmax(0,1fr));gap:12px}
.grid3{display:grid;grid-template-columns:repeat(3,minmax(0,1fr));gap:12px}
.card{background:var(--card);border:1px solid var(--border);border-radius:18px;padding:16px;box-shadow:var(--shadow)}
.metric{min-height:105px}.metric .v{font-size:25px;font-weight:900;margin-top:8px}.muted{color:var(--muted);font-size:13px}
h1{font-size:25px;margin:5px 0 18px}h2{font-size:19px;margin:5px 0 14px}
.vehicle{display:flex;gap:12px;align-items:center}.dot{width:17px;height:17px;border-radius:50%}.red{background:#ef4444}.white{background:#d1d5db;border:1px solid #9ca3af}
.list{display:flex;flex-direction:column;gap:10px}.row{display:flex;align-items:center;justify-content:space-between;gap:10px;border-bottom:1px solid #f0f1f3;padding:10px 0}.row:last-child{border-bottom:0}
.btn{border:0;border-radius:12px;padding:11px 14px;background:var(--primary);color:#fff;font-weight:700}.btn.secondary{background:#eef2f6;color:#111827}.btn.green{background:#067647}.btn.redbtn{background:#b42318}.btn.blue{background:#175cd3}
.form{display:grid;gap:10px}.form input,.form select,.form textarea{width:100%;border:1px solid var(--border);background:#fff;border-radius:12px;padding:12px;outline:none}.form textarea{min-height:100px}
.bottom{position:fixed;bottom:0;left:0;right:0;background:#fff;border-top:1px solid var(--border);height:72px;display:grid;grid-template-columns:repeat(5,1fr);z-index:20;padding-bottom:env(safe-area-inset-bottom)}
.navbtn{border:0;background:#fff;color:#667085;font-size:11px}.navbtn.active{color:#111827;font-weight:800}.navbtn span{display:block;font-size:22px;margin-bottom:2px}
.tablewrap{overflow:auto}.table{width:100%;border-collapse:collapse;min-width:650px}.table th,.table td{padding:10px;border-bottom:1px solid var(--border);text-align:right}
.badge{display:inline-block;padding:5px 9px;border-radius:999px;background:#eef2f6;font-size:12px}
.chat{height:350px;overflow:auto;background:#f8fafc;border-radius:16px;padding:12px}.msg{max-width:88%;padding:10px 12px;border-radius:14px;margin:7px 0;white-space:pre-wrap}.msg.me{margin-right:auto;background:#111827;color:#fff}.msg.ai{margin-left:auto;background:#e8eef8}
.drop{border:2px dashed #cbd5e1;border-radius:18px;padding:30px;text-align:center;background:#fff}
.small{font-size:12px}
@media(max-width:650px){.grid3{grid-template-columns:1fr 1fr}.card{padding:13px}main{padding:12px}.top{padding:12px}.metric .v{font-size:21px}}
</style>
</head>
<body>
<header class="top">
  <div><div class="brand">🍦 رمضان جاد / ظاظا</div><div class="user" id="userInfo">مدير النظام</div></div>
  <button class="btn secondary" onclick="logout()">خروج</button>
</header>

<main>
<section id="home" class="page active">
<h1>الرئيسية</h1>
<div class="grid">
  <div class="card metric"><div class="muted">مبيعات اليوم</div><div class="v" id="mToday">0 ج.م</div></div>
  <div class="card metric"><div class="muted">مبيعات الشهر</div><div class="v" id="mMonth">0 ج.م</div></div>
  <div class="card metric"><div class="muted">عدد الفواتير</div><div class="v" id="mInvoices">0</div></div>
  <div class="card metric"><div class="muted">قيمة المخزون</div><div class="v" id="mStock">0 ج.م</div></div>
</div>
<div class="grid" style="margin-top:12px">
  <div class="card"><h2>🚚 العربيات</h2><div id="vehicles" class="list"></div></div>
  <div class="card"><h2>📦 تنبيهات المخزون</h2><div id="lowStock" class="list"></div></div>
</div>
<div class="card" style="margin-top:12px"><h2>🔥 الأكثر مبيعًا</h2><div id="topProducts" class="list"></div></div>
</section>

<section id="upload" class="page">
<h1>📸 رفع فاتورة بالذكاء الاصطناعي</h1>
<div class="card">
<p class="muted">ارفع صورة الفاتورة وسيتم استخراج الأصناف والكميات والأسعار ثم حفظها في قاعدة البيانات.</p>
<form id="uploadForm" class="form">
<input type="file" id="invoiceImage" accept="image/*" required>
<input id="invoiceCustomer" placeholder="اسم العميل (اختياري)">
<button class="btn blue" type="submit">تحليل الفاتورة وحفظها</button>
</form>
<div id="uploadResult" style="margin-top:12px"></div>
</div>
</section>

<section id="invoices" class="page">
<h1>🧾 الفواتير</h1>
<div class="card">
<div style="display:flex;gap:8px;flex-wrap:wrap;margin-bottom:12px">
<button class="btn green" onclick="exportExcel()">📊 تصدير Excel</button>
<button class="btn secondary" onclick="loadInvoices()">تحديث</button>
</div>
<div class="tablewrap"><table class="table"><thead><tr><th>رقم</th><th>العميل</th><th>المنتجات</th><th>الإجمالي</th><th>التاريخ</th></tr></thead><tbody id="invoiceRows"></tbody></table></div>
</div>
</section>

<section id="link" class="page">
<h1>🔗 الربط</h1>
<div class="card">
<h2>Telegram</h2>
<form class="form" id="telegramForm">
<input id="tgToken" type="password" placeholder="Bot Token">
<input id="tgChat" placeholder="Chat ID">
<button class="btn blue">حفظ إعدادات Telegram</button>
</form>
<p id="tgStatus" class="muted"></p>
</div>
<div class="card" style="margin-top:12px">
<h2>Facebook / Messenger</h2>
<p class="muted">نقاط Webhook جاهزة في النظام: <code>/webhook/facebook</code></p>
<p class="muted">ضع Verify Token و Page Access Token في إعدادات الخادم.</p>
</div>
<div class="card" style="margin-top:12px">
<h2>WhatsApp</h2>
<p class="muted">نقطة Webhook جاهزة: <code>/webhook/whatsapp</code></p>
<p class="muted">التكامل الفعلي يحتاج بيانات Meta Business الخاصة بحسابك.</p>
</div>
</section>

<section id="settings" class="page">
<h1>⚙️ الإعدادات</h1>
<div class="card">
<h2>معلومات النظام</h2>
<div class="list">
<div class="row"><span>الحساب الحالي</span><b id="settingsUser"></b></div>
<div class="row"><span>الدور</span><b>مدير</b></div>
<div class="row"><span>قاعدة البيانات</span><b>SQLAlchemy</b></div>
</div>
</div>
<div class="card" style="margin-top:12px">
<h2>🤖 مساعد Z</h2>
<div id="chat" class="chat"><div class="msg ai">أهلًا بك. اسألني عن المبيعات أو الفواتير أو المخزون.</div></div>
<form id="chatForm" style="display:flex;gap:8px;margin-top:10px">
<input id="chatInput" style="flex:1;border:1px solid #ddd;border-radius:12px;padding:12px" placeholder="مثال: كام مبيعات العربية الحمراء النهارده؟">
<button class="btn">إرسال</button>
</form>
</div>
</section>
</main>

<nav class="bottom">
<button class="navbtn active" data-page="home"><span>🏠</span>الرئيسية</button>
<button class="navbtn" data-page="upload"><span>📸</span>رفع فاتورة</button>
<button class="navbtn" data-page="invoices"><span>🧾</span>الفواتير</button>
<button class="navbtn" data-page="link"><span>🔗</span>الربط</button>
<button class="navbtn" data-page="settings"><span>⚙️</span>الإعدادات</button>
</nav>

<script>
const money = n => new Intl.NumberFormat('ar-EG',{maximumFractionDigits:2}).format(Number(n||0))+' ج.م';
async function api(url, options={}) {
  const r = await fetch(url, options);
  if (r.status === 401) { location.href='/login'; throw new Error('غير مسجل'); }
  const data = await r.json().catch(()=>({}));
  if (!r.ok) throw new Error(data.detail || data.error || 'حدث خطأ');
  return data;
}
function showPage(id){
  document.querySelectorAll('.page').forEach(x=>x.classList.remove('active'));
  document.getElementById(id).classList.add('active');
  document.querySelectorAll('.navbtn').forEach(x=>x.classList.toggle('active',x.dataset.page===id));
  if(id==='home') loadDashboard();
  if(id==='invoices') loadInvoices();
  if(id==='link') loadTelegram();
}
document.querySelectorAll('.navbtn').forEach(x=>x.addEventListener('click',()=>showPage(x.dataset.page)));

async function loadDashboard(){
  try{
    const d=await api('/api/dashboard');
    document.getElementById('mToday').textContent=money(d.today_sales);
    document.getElementById('mMonth').textContent=money(d.month_sales);
    document.getElementById('mInvoices').textContent=d.today_invoices;
    document.getElementById('mStock').textContent=money(d.stock_value);
    document.getElementById('vehicles').innerHTML=d.vehicles.map(v=>`
      <div class="row"><div class="vehicle"><i class="dot ${v.color==='red'?'red':'white'}"></i><div><b>${v.name}</b><div class="muted">مخزون: ${money(v.stock_value)}</div></div></div><span class="badge">${money(v.today_sales)}</span></div>`).join('') || '<div class="muted">لا توجد عربيات</div>';
    document.getElementById('lowStock').innerHTML=d.low_stock.map(p=>`<div class="row"><span>${p.name}</span><span class="badge">${p.quantity} / حد ${p.min_quantity}</span></div>`).join('') || '<div class="muted">المخزون جيد</div>';
    document.getElementById('topProducts').innerHTML=d.top_products.map(p=>`<div class="row"><span>${p.name}</span><b>${p.quantity}</b></div>`).join('') || '<div class="muted">لا توجد مبيعات بعد</div>';
  }catch(e){console.error(e)}
}
async function loadInvoices(){
  try{
    const d=await api('/api/invoices');
    document.getElementById('invoiceRows').innerHTML=d.invoices.map(i=>`
      <tr><td>#${i.id}</td><td>${escapeHtml(i.customer_name)}</td><td>${escapeHtml(i.items.map(x=>`${x.name||'صنف'} × ${x.quantity||0}`).join('، '))}</td><td>${money(i.total_amount)}</td><td>${new Date(i.created_at).toLocaleString('ar-EG')}</td></tr>`).join('');
  }catch(e){alert(e.message)}
}
function escapeHtml(s){return String(s??'').replace(/[&<>"']/g,m=>({'&':'&amp;','<':'&lt;','>':'&gt;','"':'&quot;',"'":'&#039;'}[m]))}
function exportExcel(){location.href='/api/export-excel'}
document.getElementById('uploadForm').addEventListener('submit',async e=>{
  e.preventDefault();
  const f=document.getElementById('invoiceImage').files[0];
  if(!f)return;
  if(f.size>10*1024*1024){alert('الصورة أكبر من 10 ميجابايت');return}
  const fd=new FormData();
  fd.append('file',f);
  fd.append('customer_name',document.getElementById('invoiceCustomer').value||'عميل من صورة');
  document.getElementById('uploadResult').textContent='جاري تحليل الصورة...';
  try{
    const d=await api('/api/upload-invoice',{method:'POST',body:fd});
    document.getElementById('uploadResult').innerHTML=`<div class="card"><b>تم الحفظ بنجاح — فاتورة #${d.invoice.id}</b><br>الإجمالي: ${money(d.invoice.total_amount)}<br><pre>${escapeHtml(JSON.stringify(d.invoice.items,null,2))}</pre></div>`;
    loadDashboard();
  }catch(err){document.getElementById('uploadResult').textContent='خطأ: '+err.message}
});
async function loadTelegram(){
  try{
    const d=await api('/api/integrations/telegram');
    document.getElementById('tgToken').value=d.config.bot_token||'';
    document.getElementById('tgChat').value=d.config.chat_id||'';
    document.getElementById('tgStatus').textContent=d.enabled?'Telegram مفعل':'Telegram غير مفعل';
  }catch(e){}
}
document.getElementById('telegramForm').addEventListener('submit',async e=>{
  e.preventDefault();
  try{
    const d=await api('/api/integrations/telegram',{method:'POST',headers:{'Content-Type':'application/json'},body:JSON.stringify({bot_token:document.getElementById('tgToken').value,chat_id:document.getElementById('tgChat').value})});
    document.getElementById('tgStatus').textContent=d.enabled?'تم الحفظ والتفعيل':'تم الحفظ';
  }catch(e){alert(e.message)}
});
document.getElementById('chatForm').addEventListener('submit',async e=>{
  e.preventDefault();
  const input=document.getElementById('chatInput');
  const q=input.value.trim(); if(!q)return;
  const chat=document.getElementById('chat');
  chat.innerHTML+=`<div class="msg me">${escapeHtml(q)}</div>`;
  input.value='';
  chat.scrollTop=chat.scrollHeight;
  try{
    const d=await api('/api/assistant-chat',{method:'POST',headers:{'Content-Type':'application/json'},body:JSON.stringify({message:q})});
    chat.innerHTML+=`<div class="msg ai">${escapeHtml(d.answer)}</div>`;
  }catch(e){chat.innerHTML+=`<div class="msg ai">تعذر تنفيذ الطلب: ${escapeHtml(e.message)}</div>`}
  chat.scrollTop=chat.scrollHeight;
});
async function logout(){
  await fetch('/logout',{method:'POST'});
  location.href='/login';
}
loadDashboard();
</script>
</body>
</html>
"""
    return HTMLResponse(html)


# -----------------------------------------------------------------------------
# Dashboard API
# -----------------------------------------------------------------------------

@app.get("/api/dashboard")
def dashboard(user: User = Depends(require_admin), db: Session = Depends(get_db)):
    now = datetime.utcnow()
    start_day = datetime(now.year, now.month, now.day)
    start_month = datetime(now.year, now.month, 1)

    today_sales = db.query(func.coalesce(func.sum(Sale.total), 0)).filter(Sale.created_at >= start_day).scalar() or 0
    month_sales = db.query(func.coalesce(func.sum(Sale.total), 0)).filter(Sale.created_at >= start_month).scalar() or 0
    today_invoices = db.query(func.count(Invoice.id)).filter(Invoice.created_at >= start_day).scalar() or 0

    products = db.query(Product).all()
    stock_value = sum((p.quantity or 0) * (p.purchase_price or 0) for p in products)

    vehicles = []
    for v in db.query(Vehicle).filter(Vehicle.is_active.is_(True)).all():
        inventories = db.query(VehicleInventory).filter(VehicleInventory.vehicle_id == v.id).all()
        value = sum((i.quantity or 0) * (i.product.purchase_price or 0) for i in inventories if i.product)
        sales = db.query(func.coalesce(func.sum(Sale.total), 0)).filter(
            Sale.vehicle_id == v.id,
            Sale.created_at >= start_day,
        ).scalar() or 0
        vehicles.append({
            "id": v.id,
            "name": v.name,
            "color": v.color,
            "driver_name": v.driver_name,
            "stock_value": value,
            "today_sales": sales,
        })

    low_stock = [
        {"id": p.id, "name": p.name, "quantity": p.quantity, "min_quantity": p.min_quantity}
        for p in products
        if p.quantity <= p.min_quantity
    ]

    rows = (
        db.query(
            Product.name,
            func.sum(SaleItem.quantity).label("quantity"),
        )
        .join(SaleItem, SaleItem.product_id == Product.id)
        .join(Sale, Sale.id == SaleItem.sale_id)
        .filter(Sale.created_at >= start_month)
        .group_by(Product.id, Product.name)
        .order_by(func.sum(SaleItem.quantity).desc())
        .limit(5)
        .all()
    )

    return {
        "today_sales": float(today_sales),
        "month_sales": float(month_sales),
        "today_invoices": int(today_invoices),
        "stock_value": float(stock_value),
        "vehicles": vehicles,
        "low_stock": low_stock,
        "top_products": [{"name": r.name, "quantity": float(r.quantity or 0)} for r in rows],
    }


# -----------------------------------------------------------------------------
# Invoice upload / Gemini
# -----------------------------------------------------------------------------

def parse_gemini_json(text_value: str) -> dict:
    cleaned = text_value.strip()

    if cleaned.startswith("```"):
        cleaned = cleaned.replace("```json", "", 1).replace("```", "").strip()

    try:
        return json.loads(cleaned)
    except json.JSONDecodeError:
        start = cleaned.find("{")
        end = cleaned.rfind("}")
        if start >= 0 and end > start:
            return json.loads(cleaned[start:end + 1])
        raise


def analyze_invoice_with_gemini(image_bytes: bytes, mime_type: str) -> dict:
    if not GEMINI_API_KEY:
        raise RuntimeError("GEMINI_API_KEY غير مضبوط في متغيرات البيئة")

    if genai is None:
        raise RuntimeError("حزمة google-genai غير مثبتة. ثبّت google-genai.")

    client = genai.Client(api_key=GEMINI_API_KEY)

    prompt = """
أنت محاسب متخصص في قراءة فواتير المنتجات.
حلل صورة الفاتورة بدقة شديدة واستخرج البيانات الظاهرة فقط.
لا تخترع أي منتج أو سعر أو كمية غير موجودة.
أعد JSON صالحًا فقط بالشكل التالي:
{
  "customer_name": "string",
  "items": [
    {
      "name": "string",
      "code": "string",
      "quantity": 0,
      "unit_price": 0,
      "total": 0
    }
  ],
  "subtotal": 0,
  "discount": 0,
  "total_amount": 0,
  "currency": "EGP",
  "confidence": 0
}
إذا لم تستطع قراءة قيمة، ضع null بدل التخمين.
confidence رقم من 0 إلى 1.
"""

    response = client.models.generate_content(
        model=GEMINI_MODEL,
        contents=[
            types.Part.from_bytes(data=image_bytes, mime_type=mime_type),
            prompt,
        ],
        config=types.GenerateContentConfig(
            temperature=0,
            response_mime_type="application/json",
        ),
    )

    return parse_gemini_json(response.text)


@app.post("/api/upload-invoice")
async def upload_invoice(
    request: Request,
    file: UploadFile = File(...),
    customer_name: str = Form(""),
    user: User = Depends(require_admin),
    db: Session = Depends(get_db),
):
    if not file.content_type or not file.content_type.startswith("image/"):
        raise HTTPException(status_code=400, detail="الملف يجب أن يكون صورة")

    data = await file.read()
    if len(data) > MAX_UPLOAD_MB * 1024 * 1024:
        raise HTTPException(status_code=413, detail=f"حجم الصورة أكبر من {MAX_UPLOAD_MB} ميجابايت")

    try:
        extracted = analyze_invoice_with_gemini(data, file.content_type)
    except Exception as exc:
        logger.exception("Gemini invoice analysis failed.")
        raise HTTPException(status_code=502, detail=f"تعذر تحليل الفاتورة بالذكاء الاصطناعي: {exc}")

    items = extracted.get("items") or []
    normalized_items = []

    for item in items:
        normalized_items.append({
            "name": str(item.get("name") or "صنف غير محدد"),
            "code": item.get("code"),
            "quantity": safe_float(item.get("quantity"), 0),
            "unit_price": safe_float(item.get("unit_price"), 0),
            "total": safe_float(item.get("total"), 0),
        })

    total = safe_float(extracted.get("total_amount"), 0)
    if total <= 0:
        total = sum(x["total"] for x in normalized_items)
        if total <= 0:
            total = sum(x["quantity"] * x["unit_price"] for x in normalized_items)

    invoice = Invoice(
        customer_name=customer_name.strip() or str(extracted.get("customer_name") or "عميل من صورة"),
        items_json=json.dumps(normalized_items, ensure_ascii=False),
        total_amount=total,
    )
    db.add(invoice)
    db.flush()

    audit(
        db,
        user.username,
        "create",
        "invoice",
        invoice.id,
        new_data=serialize_invoice(invoice),
    )

    db.commit()
    db.refresh(invoice)

    telegram_text = (
        f"🍦 فاتورة جديدة #{invoice.id}\n"
        f"العميل: {invoice.customer_name}\n"
        f"الإجمالي: {invoice.total_amount:.2f} ج.م\n"
        f"عدد الأصناف: {len(normalized_items)}"
    )
    telegram_sent = send_telegram_message(telegram_text, db)

    return {
        "ok": True,
        "invoice": serialize_invoice(invoice),
        "ai_confidence": extracted.get("confidence"),
        "telegram_sent": telegram_sent,
    }


# -----------------------------------------------------------------------------
# Invoice listing
# -----------------------------------------------------------------------------

@app.get("/api/invoices")
def list_invoices(
    limit: int = 100,
    user: User = Depends(require_admin),
    db: Session = Depends(get_db),
):
    limit = max(1, min(limit, 500))
    invoices = db.query(Invoice).order_by(Invoice.created_at.desc()).limit(limit).all()
    return {"invoices": [serialize_invoice(x) for x in invoices]}


# -----------------------------------------------------------------------------
# Excel export
# -----------------------------------------------------------------------------

@app.get("/api/export-excel")
def export_excel(user: User = Depends(require_admin), db: Session = Depends(get_db)):
    invoices = db.query(Invoice).order_by(Invoice.created_at.desc()).all()

    wb = Workbook()
    ws = wb.active
    ws.title = "الفواتير"
    ws.sheet_view.rightToLeft = True

    headers = ["رقم الفاتورة", "العميل", "المنتجات", "الإجمالي", "التاريخ"]
    ws.append(headers)

    header_fill = PatternFill("solid", fgColor="111827")
    header_font = Font(color="FFFFFF", bold=True)
    thin = Side(style="thin", color="D0D5DD")

    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = Border(bottom=thin)

    for invoice in invoices:
        items = json_loads_safe(invoice.items_json, [])
        item_text = "، ".join(
            f"{i.get('name', 'صنف')} × {i.get('quantity', 0)} × {i.get('unit_price', 0)}"
            for i in items
        )

        ws.append([
            invoice.id,
            invoice.customer_name,
            item_text,
            invoice.total_amount,
            invoice.created_at.strftime("%Y-%m-%d %H:%M") if invoice.created_at else "",
        ])

    widths = [16, 28, 70, 18, 24]
    for index, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(index)].width = width

    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = Border(bottom=thin)

    ws.freeze_panes = "A2"

    stream = io.BytesIO()
    wb.save(stream)
    stream.seek(0)

    return StreamingResponse(
        stream,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={
            "Content-Disposition": "attachment; filename=ramadan_invoices.xlsx"
        },
    )


# -----------------------------------------------------------------------------
# Integration configuration
# -----------------------------------------------------------------------------

@app.get("/api/integrations/telegram")
def get_telegram_integration(user: User = Depends(require_admin), db: Session = Depends(get_db)):
    config = telegram_config(db)
    enabled = bool((config.get("bot_token") or TELEGRAM_BOT_TOKEN) and (config.get("chat_id") or TELEGRAM_CHAT_ID))
    return {
        "platform": "telegram",
        "enabled": enabled,
        "config": {
            "bot_token": config.get("bot_token") or TELEGRAM_BOT_TOKEN,
            "chat_id": config.get("chat_id") or TELEGRAM_CHAT_ID,
        },
    }


@app.post("/api/integrations/telegram")
async def save_telegram_integration(
    request: Request,
    user: User = Depends(require_admin),
    db: Session = Depends(get_db),
):
    body = await request.json()
    config = {
        "bot_token": str(body.get("bot_token") or "").strip(),
        "chat_id": str(body.get("chat_id") or "").strip(),
    }

    row = db.query(IntegrationConfig).filter(IntegrationConfig.platform == "telegram").first()
    if not row:
        row = IntegrationConfig(platform="telegram", config_json=json.dumps(config))
        db.add(row)
    else:
        row.config_json = json.dumps(config)

    audit(db, user.username, "update", "integration_configs", "telegram", new_data=config)
    db.commit()

    enabled = bool(config["bot_token"] and config["chat_id"])
    return {"ok": True, "enabled": enabled}


# -----------------------------------------------------------------------------
# Facebook webhook
# -----------------------------------------------------------------------------

@app.get("/webhook/facebook")
def facebook_webhook_verify(request: Request):
    params = request.query_params
    mode = params.get("hub.mode")
    verify_token = params.get("hub.verify_token")
    challenge = params.get("hub.challenge")

    if mode == "subscribe" and verify_token and verify_token == FACEBOOK_VERIFY_TOKEN:
        return int(challenge or 0)

    raise HTTPException(status_code=403, detail="Facebook verification failed")


@app.post("/webhook/facebook")
async def facebook_webhook(request: Request, db: Session = Depends(get_db)):
    payload = await request.json()
    logger.info("Facebook webhook received: %s", json.dumps(payload, ensure_ascii=False))
    db.add(Notification(
        title="Facebook Webhook",
        body=json.dumps(payload, ensure_ascii=False),
        level="info",
    ))
    db.commit()
    return {"ok": True}


# -----------------------------------------------------------------------------
# WhatsApp webhook
# -----------------------------------------------------------------------------

@app.get("/webhook/whatsapp")
def whatsapp_webhook_verify(request: Request):
    params = request.query_params
    mode = params.get("hub.mode")
    verify_token = params.get("hub.verify_token")
    challenge = params.get("hub.challenge")

    if mode == "subscribe" and verify_token and verify_token == WHATSAPP_VERIFY_TOKEN:
        return int(challenge or 0)

    raise HTTPException(status_code=403, detail="WhatsApp verification failed")


@app.post("/webhook/whatsapp")
async def whatsapp_webhook(request: Request, db: Session = Depends(get_db)):
    payload = await request.json()
    logger.info("WhatsApp webhook received: %s", json.dumps(payload, ensure_ascii=False))
    db.add(Notification(
        title="WhatsApp Webhook",
        body=json.dumps(payload, ensure_ascii=False),
        level="info",
    ))
    db.commit()
    return {"ok": True}


# -----------------------------------------------------------------------------
# Telegram webhook
# -----------------------------------------------------------------------------

@app.post("/webhook/telegram")
async def telegram_webhook(request: Request, db: Session = Depends(get_db)):
    payload = await request.json()
    logger.info("Telegram webhook received: %s", json.dumps(payload, ensure_ascii=False))

    message = payload.get("message") or {}
    text_message = message.get("text") or ""

    db.add(Notification(
        title="Telegram Webhook",
        body=text_message or json.dumps(payload, ensure_ascii=False),
        level="info",
    ))
    db.commit()

    return {"ok": True}


# -----------------------------------------------------------------------------
# Smart Assistant Z
# -----------------------------------------------------------------------------

def build_sales_context(db: Session) -> dict:
    now = datetime.utcnow()
    start_day = datetime(now.year, now.month, now.day)
    start_month = datetime(now.year, now.month, 1)

    today_sales = db.query(func.coalesce(func.sum(Sale.total), 0)).filter(Sale.created_at >= start_day).scalar() or 0
    month_sales = db.query(func.coalesce(func.sum(Sale.total), 0)).filter(Sale.created_at >= start_month).scalar() or 0

    vehicle_sales = []
    for vehicle in db.query(Vehicle).filter(Vehicle.is_active.is_(True)).all():
        total = db.query(func.coalesce(func.sum(Sale.total), 0)).filter(
            Sale.vehicle_id == vehicle.id,
            Sale.created_at >= start_month,
        ).scalar() or 0
        vehicle_sales.append({
            "vehicle": vehicle.name,
            "monthly_sales": float(total),
        })

    top = (
        db.query(
            Product.name,
            func.sum(SaleItem.quantity).label("quantity"),
            func.sum(SaleItem.total).label("sales"),
        )
        .join(SaleItem, SaleItem.product_id == Product.id)
        .join(Sale, Sale.id == SaleItem.sale_id)
        .filter(Sale.created_at >= start_month)
        .group_by(Product.id, Product.name)
        .order_by(func.sum(SaleItem.total).desc())
        .limit(10)
        .all()
    )

    low_stock = [
        {"name": p.name, "quantity": p.quantity, "minimum": p.min_quantity}
        for p in db.query(Product).all()
        if p.quantity <= p.min_quantity
    ]

    return {
        "today_sales": float(today_sales),
        "month_sales": float(month_sales),
        "vehicle_sales": vehicle_sales,
        "top_products": [
            {
                "name": x.name,
                "quantity": float(x.quantity or 0),
                "sales": float(x.sales or 0),
            }
            for x in top
        ],
        "low_stock": low_stock,
    }


def deterministic_assistant_answer(question: str, context: dict) -> str:
    q = question.strip().lower()
    today = context["today_sales"]
    month = context["month_sales"]

    if "النهارده" in q or "اليوم" in q:
        return f"إجمالي مبيعات اليوم المسجلة في قاعدة البيانات: {today:,.2f} ج.م."

    if "الشهر" in q:
        return f"إجمالي مبيعات الشهر الحالي المسجلة في قاعدة البيانات: {month:,.2f} ج.م."

    if "أكتر" in q and "منتج" in q:
        if not context["top_products"]:
            return "لا توجد مبيعات مسجلة لهذا الشهر."
        lines = [
            f"{i+1}. {x['name']} — كمية {x['quantity']:,.2f} — مبيعات {x['sales']:,.2f} ج.م."
            for i, x in enumerate(context["top_products"][:5])
        ]
        return "أكثر المنتجات مبيعًا هذا الشهر:\n" + "\n".join(lines)

    if "العربية" in q and ("أعلى" in q or "مبيعات" in q):
        ordered = sorted(context["vehicle_sales"], key=lambda x: x["monthly_sales"], reverse=True)
        if not ordered:
            return "لا توجد بيانات مبيعات للعربيات."
        return "مبيعات العربيات هذا الشهر:\n" + "\n".join(
            f"- {x['vehicle']}: {x['monthly_sales']:,.2f} ج.م."
            for x in ordered
        )

    if "خلص" in q or "نفد" in q or "مخزون" in q:
        if not context["low_stock"]:
            return "لا توجد منتجات تحت الحد الأدنى للمخزون حاليًا."
        return "المنتجات التي وصلت للحد الأدنى أو أقل:\n" + "\n".join(
            f"- {x['name']}: المتاح {x['quantity']}، الحد الأدنى {x['minimum']}"
            for x in context["low_stock"]
        )

    return (
        f"بيانات النظام الحالية: مبيعات اليوم {today:,.2f} ج.م، "
        f"ومبيعات الشهر {month:,.2f} ج.م. "
        "اسألني عن مبيعات عربية، منتج، مخزون، فواتير، أو فترة محددة."
    )


@app.post("/api/assistant-chat")
async def assistant_chat(
    request: Request,
    user: User = Depends(require_admin),
    db: Session = Depends(get_db),
):
    body = await request.json()
    message = str(body.get("message") or "").strip()

    if not message:
        raise HTTPException(status_code=400, detail="اكتب سؤالك أولاً")

    context = build_sales_context(db)

    if not GEMINI_API_KEY or genai is None:
        answer = deterministic_assistant_answer(message, context)
        return {"ok": True, "answer": answer, "source": "database-rules"}

    try:
        client = genai.Client(api_key=GEMINI_API_KEY)
        system_prompt = """
أنت مساعد Z داخل نظام إدارة موزع آيس كريم.
ممنوع اختراع أي رقم.
استخدم بيانات قاعدة البيانات الموجودة في السياق فقط.
إذا لم توجد المعلومة قل إنها غير موجودة.
أجب بالعربية المصرية الواضحة والمختصرة.
لا تدّعي تنفيذ عملية في قاعدة البيانات.
"""
        prompt = (
            system_prompt
            + "\nبيانات قاعدة البيانات:\n"
            + json.dumps(context, ensure_ascii=False)
            + "\nسؤال المدير:\n"
            + message
        )

        response = client.models.generate_content(
            model=GEMINI_MODEL,
            contents=prompt,
            config=types.GenerateContentConfig(temperature=0.1),
        )
        answer = response.text.strip()
        if not answer:
            answer = deterministic_assistant_answer(message, context)

        return {"ok": True, "answer": answer, "source": "database+gemini"}
    except Exception:
        logger.exception("Assistant AI failed.")
        return {
            "ok": True,
            "answer": deterministic_assistant_answer(message, context),
            "source": "database-rules-fallback",
        }


# -----------------------------------------------------------------------------
# Product APIs
# -----------------------------------------------------------------------------

@app.get("/api/products")
def products(user: User = Depends(require_admin), db: Session = Depends(get_db)):
    rows = db.query(Product).order_by(Product.name.asc()).all()
    return {
        "products": [
            {
                "id": p.id,
                "name": p.name,
                "code": p.code,
                "category": p.category,
                "unit": p.unit,
                "purchase_price": p.purchase_price,
                "sale_price": p.sale_price,
                "quantity": p.quantity,
                "min_quantity": p.min_quantity,
                "expiry_date": p.expiry_date,
            }
            for p in rows
        ]
    }


@app.post("/api/products")
async def create_product(
    request: Request,
    user: User = Depends(require_admin),
    db: Session = Depends(get_db),
):
    body = await request.json()

    product = Product(
        name=str(body.get("name") or "").strip(),
        code=str(body.get("code") or "").strip() or None,
        category=str(body.get("category") or "").strip(),
        unit=str(body.get("unit") or "قطعة").strip(),
        purchase_price=safe_float(body.get("purchase_price")),
        sale_price=safe_float(body.get("sale_price")),
        quantity=safe_float(body.get("quantity")),
        min_quantity=safe_float(body.get("min_quantity")),
        image_url=body.get("image_url"),
        expiry_date=body.get("expiry_date"),
    )

    if not product.name:
        raise HTTPException(status_code=400, detail="اسم المنتج مطلوب")

    db.add(product)
    db.flush()
    audit(db, user.username, "create", "product", product.id, new_data=body)
    db.commit()

    return {"ok": True, "id": product.id}


# -----------------------------------------------------------------------------
# Vehicle inventory / loading
# -----------------------------------------------------------------------------

@app.get("/api/vehicles")
def vehicles(user: User = Depends(require_admin), db: Session = Depends(get_db)):
    rows = db.query(Vehicle).filter(Vehicle.is_active.is_(True)).all()
    result = []

    for v in rows:
        inv = db.query(VehicleInventory).filter(VehicleInventory.vehicle_id == v.id).all()
        result.append({
            "id": v.id,
            "name": v.name,
            "color": v.color,
            "driver_name": v.driver_name,
            "inventory": [
                {
                    "product_id": x.product_id,
                    "product_name": x.product.name if x.product else "",
                    "quantity": x.quantity,
                    "unit_price": x.product.sale_price if x.product else 0,
                }
                for x in inv
            ],
        })

    return {"vehicles": result}


@app.post("/api/load-vehicle")
async def load_vehicle(
    request: Request,
    user: User = Depends(require_admin),
    db: Session = Depends(get_db),
):
    body = await request.json()
    vehicle_id = int(body.get("vehicle_id"))
    items = body.get("items") or []

    vehicle = db.query(Vehicle).filter(Vehicle.id == vehicle_id, Vehicle.is_active.is_(True)).first()
    if not vehicle:
        raise HTTPException(status_code=404, detail="العربية غير موجودة")

    loaded = []

    try:
        for item in items:
            product_id = int(item["product_id"])
            qty = safe_float(item["quantity"])
            if qty <= 0:
                continue

            product = db.query(Product).filter(Product.id == product_id).first()
            if not product:
                raise HTTPException(status_code=404, detail=f"المنتج {product_id} غير موجود")
            if product.quantity < qty:
                raise HTTPException(
                    status_code=400,
                    detail=f"المخزون غير كافٍ للمنتج: {product.name}. المتاح {product.quantity}",
                )

            inv = db.query(VehicleInventory).filter(
                VehicleInventory.vehicle_id == vehicle_id,
                VehicleInventory.product_id == product_id,
            ).first()

            if not inv:
                inv = VehicleInventory(vehicle_id=vehicle_id, product_id=product_id, quantity=0)
                db.add(inv)

            product.quantity -= qty
            inv.quantity += qty

            movement = StockMovement(
                product_id=product_id,
                vehicle_id=vehicle_id,
                movement_type="load_to_vehicle",
                quantity=qty,
                note=f"تحميل إلى {vehicle.name}",
                created_by=user.username,
            )
            db.add(movement)
            loaded.append({"product_id": product_id, "quantity": qty})

        audit(db, user.username, "load_vehicle", "vehicle", vehicle.id, new_data=loaded)
        db.commit()
    except Exception:
        db.rollback()
        raise

    return {"ok": True, "vehicle": vehicle.name, "items": loaded}


# -----------------------------------------------------------------------------
# Sales
# -----------------------------------------------------------------------------

@app.post("/api/sales")
async def create_sale(
    request: Request,
    user: User = Depends(require_admin),
    db: Session = Depends(get_db),
):
    body = await request.json()
    vehicle_id = int(body.get("vehicle_id"))
    customer_id = body.get("customer_id")
    payment_method = str(body.get("payment_method") or "نقدي")
    discount = safe_float(body.get("discount"))
    items = body.get("items") or []

    vehicle = db.query(Vehicle).filter(Vehicle.id == vehicle_id).first()
    if not vehicle:
        raise HTTPException(status_code=404, detail="العربية غير موجودة")

    if not items:
        raise HTTPException(status_code=400, detail="لا توجد أصناف في الفاتورة")

    normalized = []
    subtotal = 0

    try:
        for raw in items:
            product_id = int(raw["product_id"])
            qty = safe_float(raw["quantity"])
            if qty <= 0:
                continue

            product = db.query(Product).filter(Product.id == product_id).first()
            if not product:
                raise HTTPException(status_code=404, detail="منتج غير موجود")

            inv = db.query(VehicleInventory).filter(
                VehicleInventory.vehicle_id == vehicle_id,
                VehicleInventory.product_id == product_id,
            ).first()

            if not inv or inv.quantity < qty:
                available = inv.quantity if inv else 0
                raise HTTPException(
                    status_code=400,
                    detail=f"الكمية غير متاحة من {product.name}. المتاح في العربية: {available}",
                )

            unit_price = safe_float(raw.get("unit_price"), product.sale_price)
            line_total = qty * unit_price
            subtotal += line_total

            normalized.append({
                "product_id": product.id,
                "name": product.name,
                "quantity": qty,
                "unit_price": unit_price,
                "purchase_price": product.purchase_price,
                "total": line_total,
            })

        total = max(0, subtotal - discount)

        sale = Sale(
            vehicle_id=vehicle_id,
            customer_id=int(customer_id) if customer_id else None,
            payment_method=payment_method,
            subtotal=subtotal,
            discount=discount,
            total=total,
            created_by=user.username,
        )
        db.add(sale)
        db.flush()

        for item in normalized:
            inv = db.query(VehicleInventory).filter(
                VehicleInventory.vehicle_id == vehicle_id,
                VehicleInventory.product_id == item["product_id"],
            ).first()
            inv.quantity -= item["quantity"]

            db.add(SaleItem(
                sale_id=sale.id,
                product_id=item["product_id"],
                quantity=item["quantity"],
                unit_price=item["unit_price"],
                purchase_price=item["purchase_price"],
                total=item["total"],
            ))

            db.add(StockMovement(
                product_id=item["product_id"],
                vehicle_id=vehicle_id,
                movement_type="sale",
                quantity=-item["quantity"],
                reference_id=str(sale.id),
                created_by=user.username,
            ))

        customer = None
        if customer_id:
            customer = db.query(Customer).filter(Customer.id == int(customer_id)).first()
            if customer:
                customer.total_purchases += total
                if payment_method == "آجل":
                    customer.balance += total
                    db.add(CustomerTransaction(
                        customer_id=customer.id,
                        transaction_type="sale_credit",
                        amount=total,
                        reference_id=str(sale.id),
                        note="بيع آجل",
                    ))

        invoice = Invoice(
            customer_name=customer.name if customer else "عميل نقدي",
            items_json=json.dumps(normalized, ensure_ascii=False),
            total_amount=total,
        )
        db.add(invoice)
        db.flush()
        sale.invoice_id = invoice.id

        audit(
            db,
            user.username,
            "create_sale",
            "sale",
            sale.id,
            new_data={
                "vehicle_id": vehicle_id,
                "items": normalized,
                "total": total,
            },
        )
        db.commit()
        db.refresh(invoice)

        send_telegram_message(
            f"🧾 فاتورة بيع #{invoice.id}\nالعربية: {vehicle.name}\nالإجمالي: {total:.2f} ج.م",
            db,
        )

        return {
            "ok": True,
            "sale_id": sale.id,
            "invoice": serialize_invoice(invoice),
        }
    except Exception:
        db.rollback()
        raise


# -----------------------------------------------------------------------------
# Returns / damages / expenses
# -----------------------------------------------------------------------------

@app.post("/api/returns")
async def create_return(
    request: Request,
    user: User = Depends(require_admin),
    db: Session = Depends(get_db),
):
    body = await request.json()
    vehicle_id = int(body["vehicle_id"])
    product_id = int(body["product_id"])
    quantity = safe_float(body["quantity"])
    reason = str(body.get("reason") or "مرتجع من العميل")

    inv = db.query(VehicleInventory).filter(
        VehicleInventory.vehicle_id == vehicle_id,
        VehicleInventory.product_id == product_id,
    ).first()

    if not inv or inv.quantity < quantity:
        raise HTTPException(status_code=400, detail="الكمية غير متاحة في العربية")

    inv.quantity -= quantity

    product = db.query(Product).filter(Product.id == product_id).first()
    product.quantity += quantity

    row = ReturnRecord(
        vehicle_id=vehicle_id,
        product_id=product_id,
        quantity=quantity,
        reason=reason,
        created_by=user.username,
    )
    db.add(row)

    db.add(StockMovement(
        product_id=product_id,
        vehicle_id=vehicle_id,
        movement_type="return_to_warehouse",
        quantity=quantity,
        reference_id=str(row.id),
        note=reason,
        created_by=user.username,
    ))

    audit(db, user.username, "return", "returns", row.id, new_data=body)
    db.commit()

    return {"ok": True, "id": row.id}


@app.post("/api/damages")
async def create_damage(
    request: Request,
    user: User = Depends(require_admin),
    db: Session = Depends(get_db),
):
    body = await request.json()
    vehicle_id = int(body["vehicle_id"])
    product_id = int(body["product_id"])
    quantity = safe_float(body["quantity"])
    reason = str(body.get("reason") or "تالف")

    inv = db.query(VehicleInventory).filter(
        VehicleInventory.vehicle_id == vehicle_id,
        VehicleInventory.product_id == product_id,
    ).first()

    if not inv or inv.quantity < quantity:
        raise HTTPException(status_code=400, detail="الكمية غير متاحة في العربية")

    inv.quantity -= quantity

    row = Damage(
        vehicle_id=vehicle_id,
        product_id=product_id,
        quantity=quantity,
        reason=reason,
        created_by=user.username,
    )
    db.add(row)

    db.add(StockMovement(
        product_id=product_id,
        vehicle_id=vehicle_id,
        movement_type="damage",
        quantity=-quantity,
        reference_id=str(row.id),
        note=reason,
        created_by=user.username,
    ))

    audit(db, user.username, "damage", "damages", row.id, new_data=body)
    db.commit()

    return {"ok": True, "id": row.id}


@app.post("/api/expenses")
async def create_expense(
    request: Request,
    user: User = Depends(require_admin),
    db: Session = Depends(get_db),
):
    body = await request.json()
    amount = safe_float(body.get("amount"))
    reason = str(body.get("reason") or "").strip()
    vehicle_id = body.get("vehicle_id")

    if amount <= 0 or not reason:
        raise HTTPException(status_code=400, detail="القيمة والسبب مطلوبان")

    row = Expense(
        vehicle_id=int(vehicle_id) if vehicle_id else None,
        amount=amount,
        reason=reason,
        created_by=user.username,
    )
    db.add(row)
    db.flush()

    audit(db, user.username, "expense", "expenses", row.id, new_data=body)
    db.commit()

    return {"ok": True, "id": row.id}


# -----------------------------------------------------------------------------
# Daily settlement
# -----------------------------------------------------------------------------

@app.post("/api/settlements/close")
async def close_settlement(
    request: Request,
    user: User = Depends(require_admin),
    db: Session = Depends(get_db),
):
    body = await request.json()
    vehicle_id = int(body["vehicle_id"])
    settlement_date = str(body.get("date") or date.today().isoformat())
    actual_cash = safe_float(body.get("actual_cash"))

    vehicle = db.query(Vehicle).filter(Vehicle.id == vehicle_id).first()
    if not vehicle:
        raise HTTPException(status_code=404, detail="العربية غير موجودة")

    existing = db.query(DailySettlement).filter(
        DailySettlement.vehicle_id == vehicle_id,
        DailySettlement.settlement_date == settlement_date,
    ).first()

    if existing:
        raise HTTPException(status_code=409, detail="تم إغلاق يوم العربية بالفعل")

    start = datetime.fromisoformat(settlement_date)
    end = start + timedelta(days=1)

    sales_value = db.query(func.coalesce(func.sum(Sale.total), 0)).filter(
        Sale.vehicle_id == vehicle_id,
        Sale.created_at >= start,
        Sale.created_at < end,
        Sale.payment_method == "نقدي",
    ).scalar() or 0

    expense_value = db.query(func.coalesce(func.sum(Expense.amount), 0)).filter(
        Expense.vehicle_id == vehicle_id,
        Expense.created_at >= start,
        Expense.created_at < end,
    ).scalar() or 0

    expected_cash = float(sales_value) - float(expense_value)
    difference = actual_cash - expected_cash

    row = DailySettlement(
        vehicle_id=vehicle_id,
        settlement_date=settlement_date,
        sales_value=float(sales_value),
        expenses_value=float(expense_value),
        expected_cash=expected_cash,
        actual_cash=actual_cash,
        difference=difference,
        closed_by=user.username,
    )
    db.add(row)

    audit(
        db,
        user.username,
        "close_day",
        "daily_settlement",
        row.id,
        new_data={
            "vehicle": vehicle.name,
            "date": settlement_date,
            "expected_cash": expected_cash,
            "actual_cash": actual_cash,
            "difference": difference,
        },
    )
    db.commit()

    return {
        "ok": True,
        "vehicle": vehicle.name,
        "expected_cash": expected_cash,
        "actual_cash": actual_cash,
        "difference": difference,
    }


# -----------------------------------------------------------------------------
# Customers / payments
# -----------------------------------------------------------------------------

@app.get("/api/customers")
def customers(user: User = Depends(require_admin), db: Session = Depends(get_db)):
    rows = db.query(Customer).order_by(Customer.name.asc()).all()
    return {
        "customers": [
            {
                "id": x.id,
                "name": x.name,
                "phone": x.phone,
                "area": x.area,
                "customer_type": x.customer_type,
                "balance": x.balance,
                "total_purchases": x.total_purchases,
            }
            for x in rows
        ]
    }


@app.post("/api/customers")
async def create_customer(
    request: Request,
    user: User = Depends(require_admin),
    db: Session = Depends(get_db),
):
    body = await request.json()
    name = str(body.get("name") or "").strip()
    if not name:
        raise HTTPException(status_code=400, detail="اسم العميل مطلوب")

    row = Customer(
        name=name,
        phone=str(body.get("phone") or ""),
        area=str(body.get("area") or ""),
        customer_type=str(body.get("customer_type") or "عميل عادي"),
    )
    db.add(row)
    db.flush()

    audit(db, user.username, "create", "customer", row.id, new_data=body)
    db.commit()

    return {"ok": True, "id": row.id}


@app.post("/api/payments")
async def create_payment(
    request: Request,
    user: User = Depends(require_admin),
    db: Session = Depends(get_db),
):
    body = await request.json()
    customer_id = int(body["customer_id"])
    amount = safe_float(body.get("amount"))
    payment_method = str(body.get("payment_method") or "نقدي")
    note = str(body.get("note") or "")

    if amount <= 0:
        raise HTTPException(status_code=400, detail="قيمة التحصيل يجب أن تكون أكبر من صفر")

    customer = db.query(Customer).filter(Customer.id == customer_id).first()
    if not customer:
        raise HTTPException(status_code=404, detail="العميل غير موجود")

    customer.balance = max(0, customer.balance - amount)

    row = Payment(
        customer_id=customer_id,
        amount=amount,
        payment_method=payment_method,
        note=note,
        created_by=user.username,
    )
    db.add(row)

    db.add(CustomerTransaction(
        customer_id=customer_id,
        transaction_type="payment",
        amount=amount,
        reference_id=str(row.id),
        note=note,
    ))

    audit(db, user.username, "payment", "customers", customer_id, new_data={
        "amount": amount,
        "balance_after": customer.balance,
    })
    db.commit()

    return {"ok": True, "id": row.id, "remaining_balance": customer.balance}


# -----------------------------------------------------------------------------
# Health / diagnostics
# -----------------------------------------------------------------------------

@app.get("/health")
def health(db: Session = Depends(get_db)):
    try:
        db.execute(text("SELECT 1"))
        return {
            "status": "ok",
            "database": "connected",
            "time": datetime.utcnow().isoformat(),
        }
    except Exception as exc:
        return JSONResponse(
            status_code=503,
            content={"status": "error", "database": str(exc)},
        )


@app.get("/api/me")
def me(request: Request, db: Session = Depends(get_db)):
    user = current_user(request, db)
    if not user:
        raise HTTPException(status_code=401, detail="غير مسجل")
    return {
        "id": user.id,
        "username": user.username,
        "role": user.role,
    }


# -----------------------------------------------------------------------------
# Application startup
# -----------------------------------------------------------------------------

@app.on_event("startup")
def startup_event():
    initialize_database()
    logger.info("Application started: %s", APP_NAME)


if __name__ == "__main__":
    import uvicorn

    uvicorn.run(
        "main:app",
        host=os.getenv("HOST", "0.0.0.0"),
        port=int(os.getenv("PORT", "8000")),
        reload=os.getenv("RELOAD", "false").lower() == "true",
    )
